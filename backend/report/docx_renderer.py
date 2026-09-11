import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from backend.report.charts import render_analysis_bars, render_defect_donut
from backend.report.schemas import ReportData

NAVY = RGBColor(0x1B, 0x3A, 0x5F)


def _style_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def render_docx(data: ReportData) -> bytes:
    doc = Document()

    title = doc.add_heading(data.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    r = sub.add_run(data.subtitle)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = NAVY

    meta = doc.add_paragraph()
    meta.add_run(
        f"Session {data.session_id}  ·  Generated {data.generated_at:%Y-%m-%d %H:%M UTC}  ·  {data.brand}"
    ).italic = True

    _style_heading(doc, "Output Summary", level=2)
    doc.add_paragraph(data.methodology_note)
    doc.add_paragraph(data.executive_summary)
    doc.add_paragraph(
        f"Severity: {data.severity}  |  Quality/Yield: {data.overall_quality_score}%  |  "
        f"Regions: {data.detection_count}  |  Vision confidence: {data.defect_confidence_pct:.0f}%"
    )

    _style_heading(doc, "Defect Distribution", level=2)
    doc.add_picture(io.BytesIO(render_defect_donut(data)), width=Inches(3.2))

    _style_heading(doc, "Analysis Statistics", level=2)
    doc.add_picture(io.BytesIO(render_analysis_bars(data)), width=Inches(5.2))

    _style_heading(doc, "Cause Analysis", level=2)
    doc.add_paragraph(data.process_insight)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Cause", "Probability", "Reasoning"
    for cause in sorted(data.causes, key=lambda c: -c.score):
        row = table.add_row().cells
        row[0].text = cause.name
        row[1].text = f"{cause.score}%"
        row[2].text = cause.explanation

    doc.add_page_break()

    _style_heading(doc, "Maintenance Items", level=2)
    for item in data.maintenance_items:
        doc.add_paragraph(item, style="List Bullet")

    _style_heading(doc, "Diagnostic Findings", level=2)
    for item in data.diagnostic_findings:
        doc.add_paragraph(item, style="List Bullet")

    if data.similar_case_note:
        _style_heading(doc, "Similar Historical Cases", level=2)
        doc.add_paragraph(data.similar_case_note)

    _style_heading(doc, "Recommended Troubleshooting Sequence", level=2)
    for i, step in enumerate(data.action_plan, 1):
        doc.add_paragraph(f"{i}. {step}")

    _style_heading(doc, "Geometric Feature Extraction", level=2)
    geo = doc.add_table(rows=1, cols=4)
    geo.style = "Light Grid Accent 1"
    gh = geo.rows[0].cells
    gh[0].text, gh[1].text, gh[2].text, gh[3].text = (
        "ID",
        "Defect Class",
        "Confidence",
        "Area (px2)",
    )
    if data.detections:
        for d in data.detections:
            row = geo.add_row().cells
            row[0].text = str(d.id)
            row[1].text = d.defect_class
            row[2].text = f"{d.confidence_pct:.1f}%"
            row[3].text = f"{d.area_px:.0f}" if d.area_px is not None else "-"
    else:
        row = geo.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = "-", "None", "100%", "0"

    _style_heading(doc, "Engineer Notes", level=2)
    doc.add_paragraph(data.engineer_notes or "-")

    footer = doc.add_paragraph()
    footer.add_run(data.footer).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
